"""Reads text from an uploaded notes file (.txt, .md, .pdf, .docx)."""
import io
import logging
import os
import re

from docx import Document
from pypdf import PdfReader

import config

logger = logging.getLogger(__name__)
ALLOWED_EXTENSIONS = {".txt", ".md", ".pdf", ".docx"}


class FileError(Exception):
    """A file problem with a message that is safe to show to the user."""

    def __init__(self, message, status=400):
        super().__init__(message)
        self.message = message
        self.status = status


def _read_text_file(data):
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")


def _read_pdf(data):
    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            opened = reader.decrypt("")   # many PDFs are "encrypted" but open without a password
        except Exception:
            raise FileError("This PDF uses encryption that could not be opened. Open it in a PDF viewer, "
                            "choose Print > Save as PDF to make a plain copy, then upload that copy.")
        if not opened:
            raise FileError("This PDF is password protected. Remove the password and try again.")
    parts, total = [], 0
    for page in reader.pages:
        try:
            page_text = page.extract_text() or ""
        except Exception:                 # skip a page that cannot be read
            page_text = ""
        parts.append(page_text)
        total += len(page_text)
        if total > config.MAX_NOTES_CHARS:   # enough text, stop reading more pages
            break
    return "\n".join(parts)


def _read_docx(data):
    document = Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def _clean(text):
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    return re.sub(r"\n{3,}", "\n\n", text).strip()


def extract_text(file):
    """Return {"title", "text", "truncated"} for an uploaded file, or raise FileError."""
    if file is None or not file.filename:
        raise FileError("Please choose a file to upload.")

    extension = os.path.splitext(file.filename)[1].lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise FileError("Unsupported file type. Upload a .txt, .md, .pdf or .docx file.")

    data = file.read()
    if not data:
        raise FileError("The file is empty.")
    if len(data) > config.MAX_UPLOAD_BYTES:
        raise FileError(f"The file is too large (maximum {config.MAX_UPLOAD_MB} MB).", 413)

    try:
        if extension == ".pdf":
            text = _read_pdf(data)
        elif extension == ".docx":
            text = _read_docx(data)
        else:
            text = _read_text_file(data)
    except FileError:
        raise
    except Exception:
        logger.error("Could not read an uploaded %s file", extension)
        raise FileError("Could not read this file. It may be damaged or password protected.")

    text = _clean(text)
    if not text:
        raise FileError("No readable text was found. Scanned PDFs and images are not supported.")

    title = os.path.splitext(os.path.basename(file.filename))[0].replace("_", " ").replace("-", " ").strip()
    return {
        "title": (title or "Uploaded notes")[:config.MAX_TITLE_CHARS],
        "text": text[:config.MAX_NOTES_CHARS],
        "truncated": len(text) > config.MAX_NOTES_CHARS,
    }
